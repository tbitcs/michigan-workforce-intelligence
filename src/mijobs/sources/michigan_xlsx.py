from __future__ import annotations

import io
import re
from collections.abc import Iterable
from datetime import date
from decimal import Decimal, InvalidOperation
from typing import ClassVar

from openpyxl import load_workbook
from openpyxl.workbook.workbook import Workbook
from openpyxl.worksheet.worksheet import Worksheet

from mijobs.domain import ObservationInput
from mijobs.sources.base import FetchedArtifact, SourceFetchError


def _norm(value: object) -> str:
    text = "" if value is None else str(value)
    return re.sub(r"[^a-z0-9]+", " ", text.casefold()).strip()


def _decimal(value: object) -> Decimal | None:
    if value is None:
        return None
    text = str(value).strip().replace(",", "").replace("$", "").replace("%", "")
    if text in {"", "*", "**", "#", "N/A", "NA"}:
        return None
    try:
        return Decimal(text)
    except InvalidOperation:
        return None


def _workbook(artifact: FetchedArtifact) -> Workbook:
    try:
        return load_workbook(io.BytesIO(artifact.content), read_only=True, data_only=True)
    except Exception as exc:  # openpyxl raises several format/zip errors
        raise SourceFetchError("invalid XLSX workbook") from exc


def _header_map(
    sheet: Worksheet,
    aliases: dict[str, set[str]],
    *,
    required: set[str],
    search_rows: int = 30,
) -> tuple[int, dict[str, int]]:
    normalized_aliases = {
        semantic: {_norm(alias) for alias in candidates}
        for semantic, candidates in aliases.items()
    }
    for row_index, row in enumerate(
        sheet.iter_rows(min_row=1, max_row=min(search_rows, sheet.max_row), values_only=True),
        start=1,
    ):
        mapped: dict[str, int] = {}
        for column_index, value in enumerate(row):
            header = _norm(value)
            for semantic, candidates in normalized_aliases.items():
                if header in candidates and semantic not in mapped:
                    mapped[semantic] = column_index
        if required <= set(mapped):
            return row_index, mapped
    raise SourceFetchError(
        f"could not find required XLSX headers: {sorted(required)} in first {search_rows} rows"
    )


def _value(row: tuple[object, ...], mapping: dict[str, int], name: str) -> object | None:
    index = mapping.get(name)
    if index is None or index >= len(row):
        return None
    return row[index]


class MCDAProjectionParser:
    """Michigan MCDA long-term occupational projection workbook parser."""

    parser_version = "mcda-occupation-projections-xlsx/1"
    ALIASES: ClassVar[dict[str, set[str]]] = {
        "code": {"occupation code", "soc code", "occ code", "code"},
        "title": {"occupation title", "occupational title", "occupation", "occ title"},
        "base": {"2024 employment", "base employment", "base year employment", "employment 2024"},
        "projected": {
            "2034 employment",
            "projected employment",
            "projected year employment",
            "employment 2034",
        },
        "openings": {
            "annual openings",
            "projected annual openings",
            "average annual openings",
            "annual job openings",
        },
        "change": {"employment change", "numeric change", "change"},
        "pct_change": {"percent change", "percentage change", "projected growth", "growth rate"},
    }

    def normalize(
        self,
        artifact: FetchedArtifact,
        *,
        base_year: int = 2024,
        projected_year: int = 2034,
        soc_version: str = "2018",
        geography_type: str = "state",
        geography_code: str = "26",
        geography_name: str = "Michigan",
        sheet_name: str | None = None,
    ) -> list[ObservationInput]:
        workbook = _workbook(artifact)
        sheet = workbook[sheet_name] if sheet_name else workbook.active
        header_row, mapping = _header_map(
            sheet,
            self.ALIASES,
            required={"code", "title", "base", "projected", "openings"},
        )
        observations: list[ObservationInput] = []
        for row in sheet.iter_rows(min_row=header_row + 1, values_only=True):
            code_raw = _value(row, mapping, "code")
            title_raw = _value(row, mapping, "title")
            if code_raw is None or title_raw is None:
                continue
            code = str(code_raw).strip()
            title = str(title_raw).strip()
            if not code or not title:
                continue
            shared = {
                "occupation_title": title,
                "base_year": base_year,
                "projected_year": projected_year,
                "source_dataset_version": artifact.dataset_version,
                "measure_kind": "official_projection",
            }
            metrics: Iterable[tuple[str, str, str, object | None, date | None, date | None, str]] = (
                (
                    "employment_base",
                    "mcda.projection.occupation.base_employment",
                    "jobs",
                    _value(row, mapping, "base"),
                    date(base_year, 1, 1),
                    date(base_year, 12, 31),
                    f"base_year_{base_year}",
                ),
                (
                    "employment_projected",
                    "mcda.projection.occupation.projected_employment",
                    "jobs",
                    _value(row, mapping, "projected"),
                    date(projected_year, 1, 1),
                    date(projected_year, 12, 31),
                    f"projected_year_{projected_year}",
                ),
                (
                    "annual_openings",
                    "mcda.projection.occupation.annual_openings",
                    "jobs_per_year",
                    _value(row, mapping, "openings"),
                    date(base_year, 1, 1),
                    date(projected_year, 12, 31),
                    f"annual_average_{base_year}_{projected_year}",
                ),
                (
                    "numeric_change",
                    "mcda.projection.occupation.numeric_change",
                    "jobs",
                    _value(row, mapping, "change"),
                    date(base_year, 1, 1),
                    date(projected_year, 12, 31),
                    f"change_{base_year}_{projected_year}",
                ),
                (
                    "percent_change",
                    "mcda.projection.occupation.percent_change",
                    "percent",
                    _value(row, mapping, "pct_change"),
                    date(base_year, 1, 1),
                    date(projected_year, 12, 31),
                    f"change_{base_year}_{projected_year}",
                ),
            )
            for suffix, metric, unit, raw, start, end, basis in metrics:
                if raw is None or str(raw).strip() == "":
                    continue
                observations.append(
                    ObservationInput(
                        observation_key=(
                            f"mcda:projection:{base_year}-{projected_year}:"
                            f"{geography_type}:{geography_code}:{code}:{suffix}"
                        ),
                        metric=metric,
                        value_text=str(raw).strip(),
                        numeric_value=_decimal(raw),
                        unit=unit,
                        geography_type=geography_type,
                        geography_code=geography_code,
                        geography_name=geography_name,
                        period_start=start,
                        period_end=end,
                        period_basis=basis,
                        taxonomy_system="SOC",
                        taxonomy_version=soc_version,
                        taxonomy_code=code,
                        release_status="published_projection",
                        metadata=shared,
                    )
                )
        return observations


class MCDAOEWSParser:
    """Michigan MCDA OEWS workbook parser for employment and wage snapshots."""

    parser_version = "mcda-oews-xlsx/1"
    ALIASES: ClassVar[dict[str, set[str]]] = {
        "code": {"occ code", "occupation code", "soc code", "code"},
        "title": {"occ title", "occupation title", "occupational title", "occupation"},
        "employment": {"tot emp", "total employment", "employment", "employment estimate"},
        "employment_prse": {"emp prse", "employment prse", "employment relative standard error"},
        "hourly_mean": {"h mean", "mean hourly wage", "hourly mean wage"},
        "hourly_median": {"h median", "median hourly wage", "hourly median wage"},
        "annual_mean": {"a mean", "mean annual wage", "annual mean wage"},
        "annual_median": {"a median", "median annual wage", "annual median wage"},
    }

    def normalize(
        self,
        artifact: FetchedArtifact,
        *,
        year: int = 2025,
        soc_version: str = "2018",
        geography_type: str = "state",
        geography_code: str = "26",
        geography_name: str = "Michigan",
        sheet_name: str | None = None,
    ) -> list[ObservationInput]:
        workbook = _workbook(artifact)
        sheet = workbook[sheet_name] if sheet_name else workbook.active
        header_row, mapping = _header_map(
            sheet,
            self.ALIASES,
            required={"code", "title", "employment"},
        )
        metric_defs = (
            ("employment", "mcda.oews.employment", "jobs"),
            ("employment_prse", "mcda.oews.employment_prse", "percent"),
            ("hourly_mean", "mcda.oews.hourly_mean_wage", "dollars_per_hour"),
            ("hourly_median", "mcda.oews.hourly_median_wage", "dollars_per_hour"),
            ("annual_mean", "mcda.oews.annual_mean_wage", "dollars_per_year"),
            ("annual_median", "mcda.oews.annual_median_wage", "dollars_per_year"),
        )
        observations: list[ObservationInput] = []
        for row in sheet.iter_rows(min_row=header_row + 1, values_only=True):
            code_raw = _value(row, mapping, "code")
            title_raw = _value(row, mapping, "title")
            if code_raw is None or title_raw is None:
                continue
            code = str(code_raw).strip()
            title = str(title_raw).strip()
            if not code or not title:
                continue
            for field, metric, unit in metric_defs:
                raw = _value(row, mapping, field)
                if raw is None or str(raw).strip() == "":
                    continue
                observations.append(
                    ObservationInput(
                        observation_key=(
                            f"mcda:oews:{year}:{geography_type}:{geography_code}:{code}:{field}"
                        ),
                        metric=metric,
                        value_text=str(raw).strip(),
                        numeric_value=_decimal(raw),
                        unit=unit,
                        geography_type=geography_type,
                        geography_code=geography_code,
                        geography_name=geography_name,
                        period_basis=f"oews_snapshot_{year}",
                        taxonomy_system="SOC",
                        taxonomy_version=soc_version,
                        taxonomy_code=code,
                        release_status="published_estimate",
                        metadata={
                            "occupation_title": title,
                            "year": year,
                            "source_dataset_version": artifact.dataset_version,
                            "methodology_caveat": (
                                "OEWS is a snapshot estimate; do not treat adjacent releases "
                                "as a simple year-over-year time series."
                            ),
                        },
                    )
                )
        return observations
